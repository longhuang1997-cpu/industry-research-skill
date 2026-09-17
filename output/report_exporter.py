"""
报告导出器 - Phase 3 任务3

功能:
1. Word导出（.docx）- 使用python-docx
2. Markdown导出（.md）- 使用markdownify
3. 统一导出接口

依赖库:
- python-docx
- markdownify
- beautifulsoup4

作者: Claude Opus 5
日期: 2026-09-17
"""

import os
import re
from pathlib import Path
from typing import Optional

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    print("[ReportExporter] ⚠️  未安装 python-docx，Word导出不可用")
    print("[ReportExporter] 安装: pip install python-docx")

try:
    from markdownify import markdownify as md
    MARKDOWN_AVAILABLE = True
except ImportError:
    MARKDOWN_AVAILABLE = False
    print("[ReportExporter] ⚠️  未安装 markdownify，Markdown导出不可用")
    print("[ReportExporter] 安装: pip install markdownify")

try:
    from bs4 import BeautifulSoup
    BS4_AVAILABLE = True
except ImportError:
    BS4_AVAILABLE = False
    print("[ReportExporter] ⚠️  未安装 beautifulsoup4，导出功能不可用")
    print("[ReportExporter] 安装: pip install beautifulsoup4")


class ReportExporter:
    """报告导出器"""

    def __init__(self):
        """初始化导出器"""
        self.supported_formats = []

        if DOCX_AVAILABLE:
            self.supported_formats.append('word')
        if MARKDOWN_AVAILABLE:
            self.supported_formats.append('markdown')

    def export_report(self,
                     html_path: str,
                     output_format: str,
                     output_path: Optional[str] = None) -> str:
        """
        导出报告

        Args:
            html_path: HTML报告路径
            output_format: 'word' or 'markdown'
            output_path: 输出路径（可选，默认自动生成）

        Returns:
            输出文件路径

        Raises:
            ValueError: 格式不支持
            FileNotFoundError: HTML文件不存在
        """
        # 验证格式
        if output_format not in ['word', 'markdown']:
            raise ValueError(f"不支持的格式: {output_format}（支持: word, markdown）")

        if output_format not in self.supported_formats:
            raise ValueError(f"{output_format}导出需要安装额外依赖")

        # 验证文件存在
        if not os.path.exists(html_path):
            raise FileNotFoundError(f"HTML文件不存在: {html_path}")

        # 读取HTML
        with open(html_path, 'r', encoding='utf-8') as f:
            html_content = f.read()

        # 导出
        if output_format == 'word':
            return self._export_word(html_content, html_path, output_path)
        else:
            return self._export_markdown(html_content, html_path, output_path)

    def _export_word(self,
                    html_content: str,
                    html_path: str,
                    output_path: Optional[str] = None) -> str:
        """
        导出为Word文档

        Args:
            html_content: HTML内容
            html_path: 原HTML文件路径
            output_path: 输出路径

        Returns:
            输出文件路径
        """
        if not DOCX_AVAILABLE or not BS4_AVAILABLE:
            raise RuntimeError("Word导出需要 python-docx 和 beautifulsoup4")

        # 生成输出路径
        if output_path is None:
            output_path = html_path.replace('.html', '.docx')

        print(f"[ReportExporter] Word导出: {html_path} → {output_path}")

        # 解析HTML
        soup = BeautifulSoup(html_content, 'html.parser')

        # 创建Word文档
        doc = Document()

        # 处理主要内容
        body = soup.find('body')
        if body:
            self._process_elements(doc, body.children)
        else:
            # 如果没有body，处理整个soup
            self._process_elements(doc, soup.children)

        # 保存
        doc.save(output_path)
        print(f"[ReportExporter] ✅ Word导出完成")

        return output_path

    def _process_elements(self, doc: Document, elements):
        """
        递归处理HTML元素

        Args:
            doc: Word文档对象
            elements: HTML元素列表
        """
        for element in elements:
            # 跳过文本节点
            if not hasattr(element, 'name'):
                continue

            tag = element.name

            # 跳过不需要的标签
            if tag in ['style', 'script', 'meta', 'link']:
                continue

            # 标题
            if tag in ['h1', 'h2', 'h3', 'h4', 'h5', 'h6']:
                level = int(tag[1])
                text = element.get_text().strip()
                if text:
                    doc.add_heading(text, level=min(level, 9))  # Word最多9级标题

            # 段落
            elif tag == 'p':
                text = element.get_text().strip()
                if text:
                    paragraph = doc.add_paragraph(text)
                    self._apply_inline_styles(paragraph, element)

            # 列表
            elif tag in ['ul', 'ol']:
                for li in element.find_all('li', recursive=False):
                    text = li.get_text().strip()
                    if text:
                        doc.add_paragraph(
                            text,
                            style='List Bullet' if tag == 'ul' else 'List Number'
                        )

            # 表格
            elif tag == 'table':
                self._add_table(doc, element)

            # 块引用
            elif tag == 'blockquote':
                text = element.get_text().strip()
                if text:
                    paragraph = doc.add_paragraph(text)
                    paragraph.paragraph_format.left_indent = Inches(0.5)

            # 容器（递归处理）
            elif tag in ['div', 'section', 'article']:
                self._process_elements(doc, element.children)

    def _apply_inline_styles(self, paragraph, element):
        """
        应用内联样式（加粗、斜体）

        Args:
            paragraph: Word段落对象
            element: HTML元素
        """
        # 查找加粗
        if element.find('strong') or element.find('b'):
            for run in paragraph.runs:
                run.bold = True

        # 查找斜体
        if element.find('em') or element.find('i'):
            for run in paragraph.runs:
                run.italic = True

    def _add_table(self, doc: Document, table_element):
        """
        添加表格到Word文档

        Args:
            doc: Word文档对象
            table_element: HTML表格元素
        """
        rows = table_element.find_all('tr')
        if not rows:
            return

        # 计算列数
        first_row_cells = rows[0].find_all(['th', 'td'])
        if not first_row_cells:
            return
        cols = len(first_row_cells)

        # 创建表格
        table = doc.add_table(rows=len(rows), cols=cols)
        table.style = 'Light Grid Accent 1'

        # 填充数据
        for i, row in enumerate(rows):
            cells = row.find_all(['th', 'td'])
            for j, cell in enumerate(cells):
                if j < cols:  # 防止列数不一致
                    table.rows[i].cells[j].text = cell.get_text().strip()

    def _export_markdown(self,
                        html_content: str,
                        html_path: str,
                        output_path: Optional[str] = None) -> str:
        """
        导出为Markdown

        Args:
            html_content: HTML内容
            html_path: 原HTML文件路径
            output_path: 输出路径

        Returns:
            输出文件路径
        """
        if not MARKDOWN_AVAILABLE or not BS4_AVAILABLE:
            raise RuntimeError("Markdown导出需要 markdownify 和 beautifulsoup4")

        # 生成输出路径
        if output_path is None:
            output_path = html_path.replace('.html', '.md')

        print(f"[ReportExporter] Markdown导出: {html_path} → {output_path}")

        # 预处理HTML
        soup = BeautifulSoup(html_content, 'html.parser')

        # 移除style和script
        for tag in soup.find_all(['style', 'script']):
            tag.decompose()

        # 转换为Markdown
        markdown_content = md(
            str(soup),
            heading_style="ATX",  # 使用 # 风格的标题
            bullets="-",  # 使用 - 作为列表符号
            strip=['img']  # 暂时跳过图片（后续可优化）
        )

        # 后处理
        # 1. 清理多余换行（3个以上换行 → 2个）
        markdown_content = re.sub(r'\n{3,}', '\n\n', markdown_content)

        # 2. 清理行首空格
        markdown_content = re.sub(r'^\s+', '', markdown_content, flags=re.MULTILINE)

        # 3. 修复列表缩进
        markdown_content = re.sub(r'^  -', '-', markdown_content, flags=re.MULTILINE)

        # 保存
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(markdown_content)

        print(f"[ReportExporter] ✅ Markdown导出完成")

        return output_path


# ==================== 测试代码 ====================

if __name__ == '__main__':
    print("=" * 60)
    print("报告导出器 - 测试")
    print("=" * 60)

    # 检查依赖
    print("\n依赖检查:")
    print(f"  - python-docx: {'✅ 已安装' if DOCX_AVAILABLE else '❌ 未安装'}")
    print(f"  - markdownify: {'✅ 已安装' if MARKDOWN_AVAILABLE else '❌ 未安装'}")
    print(f"  - beautifulsoup4: {'✅ 已安装' if BS4_AVAILABLE else '❌ 未安装'}")

    # 创建测试HTML
    test_html = """
    <!DOCTYPE html>
    <html>
    <head><title>测试报告</title></head>
    <body>
        <h1>行业研究报告</h1>
        <h2>第一章：行业概述</h2>
        <p>这是一个<strong>测试段落</strong>，包含<em>斜体文本</em>。</p>

        <h3>1.1 市场规模</h3>
        <p>市场规模达到100亿元。</p>

        <ul>
            <li>特点1</li>
            <li>特点2</li>
            <li>特点3</li>
        </ul>

        <table>
            <tr><th>指标</th><th>数值</th></tr>
            <tr><td>增长率</td><td>30%</td></tr>
            <tr><td>市场份额</td><td>15%</td></tr>
        </table>
    </body>
    </html>
    """

    # 写入测试文件
    test_path = "test_report.html"
    with open(test_path, 'w', encoding='utf-8') as f:
        f.write(test_html)

    print(f"\n测试文件创建: {test_path}")

    # 测试导出
    exporter = ReportExporter()

    print(f"\n支持的格式: {exporter.supported_formats}")

    # 测试Word导出
    if 'word' in exporter.supported_formats:
        try:
            word_path = exporter.export_report(test_path, 'word')
            print(f"\n✅ Word导出成功: {word_path}")
        except Exception as e:
            print(f"\n❌ Word导出失败: {e}")

    # 测试Markdown导出
    if 'markdown' in exporter.supported_formats:
        try:
            md_path = exporter.export_report(test_path, 'markdown')
            print(f"✅ Markdown导出成功: {md_path}")

            # 显示内容
            with open(md_path, 'r', encoding='utf-8') as f:
                print("\nMarkdown内容预览:")
                print("-" * 60)
                print(f.read()[:500])
                print("-" * 60)
        except Exception as e:
            print(f"❌ Markdown导出失败: {e}")

    print("\n" + "=" * 60)
    print("✅ 测试完成")
    print("=" * 60)

    # 清理测试文件
    print("\n提示: 测试文件已生成，请手动删除:")
    print(f"  - {test_path}")
    if 'word' in exporter.supported_formats:
        print(f"  - test_report.docx")
    if 'markdown' in exporter.supported_formats:
        print(f"  - test_report.md")
